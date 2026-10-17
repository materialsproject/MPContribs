"""Excel overview of which MXene structures are present, broken or missing.

The main sheet is a grid with one row per M-X system and one column per
structure that can exist for it, grouped by thickness n, then termination,
then stacking label. Each cell shows

- ``✓`` a CONTCAR that passed every check (several configurations of the
  same composition and label count as one cell),
- ``✗`` a CONTCAR that failed (hover for the reason; details on "Problems"),
- ``○`` no CONTCAR for a structure that is in scope,

and combinations excluded from the scope are shaded grey. Rows are created
for every single transition metal found in the data, with C and with N, in
order of atomic number, followed by every other system found (double-metal,
high-entropy and carbonitride MXenes). Columns cover n = 1-3 and the
terminations none, F and O (`DEFAULT_THICKNESSES`, `DEFAULT_TERMINATIONS`),
extended by any other n or termination (including mixed terminations) found
in the data. All combinations are in scope unless excluded with
``out_of_scope`` (on the command line: ``--out-of-scope Hf:O``).

Generate it with the dataset check::

    python -m mpcontribs.lux.projects.two_d_mxenes.pipelines.build_contributions DATA_ROOT --overview overview.xlsx
"""

from __future__ import annotations

from collections import Counter
from collections.abc import Collection
from datetime import datetime
from itertools import pairwise
from pathlib import Path
from typing import get_args

from mpcontribs.lux.projects.two_d_mxenes.pipelines.build_contributions import (
    CheckRecord,
    CheckResult,
)
from mpcontribs.lux.projects.two_d_mxenes.schemas.labels import (
    N1_STACKINGS,
    TERMINATIONS,
    THICK_STACKINGS,
    NonMetal,
    Termination,
    Thickness,
    join_elements,
    parse_folder_label,
    split_elements,
)
from pymatgen.core import Element

DEFAULT_THICKNESSES: tuple[int, ...] = (1, 2, 3)
"""Thickness indices shown in the grid even when no file has them."""

DEFAULT_TERMINATIONS: tuple[str | None, ...] = (None, "F", "O")
"""Terminations (None = pristine) shown in the grid even when no file has them."""

NONMETALS: tuple[str, ...] = get_args(NonMetal)
THICKNESSES: tuple[int, ...] = get_args(Thickness)

_TERMINATION_ORDER = {t: i for i, t in enumerate(get_args(Termination))}
_NONMETAL_ORDER = {"C": 0, "N": 1, "C-N": 2}
_STACKING_ORDER = ("t", "h", "h1a", "h1b", "h2")
_SUBSCRIPTS = str.maketrans("0123456789", "₀₁₂₃₄₅₆₇₈₉")

ScopeExclusion = tuple[str, str | None]
"""(metals, terminations) to exclude, e.g. ("Hf", "O") or ("Mo-Ti", None);
metals `*` means every system, terminations None means pristine."""

GridKey = tuple[str, str, int, str | None, str]
"""(metals, nonmetals, n, terminations, label) of one grid cell; the label
has no configuration suffix."""

GOOD, BAD, MISSING, OUT_OF_SCOPE = "✓", "✗", "○", ""


def _stackings(n: int) -> list[str]:
    allowed = N1_STACKINGS if n == 1 else THICK_STACKINGS
    return [s for s in _STACKING_ORDER if s in allowed]


def _n_name(n: int) -> str:
    formula = f"M{n + 1}X{n if n > 1 else ''}".translate(_SUBSCRIPTS)
    return f"n = {n}  ({formula})"


def _t_name(term: str | None) -> str:
    return "pristine" if term is None else f"{term}-terminated"


def _z_key(elements: str) -> tuple[int, ...]:
    return tuple(Element(s).Z for s in split_elements(elements))


def _termination_key(term: str | None) -> tuple:
    if term is None:
        return (0,)
    parts = split_elements(term)
    if len(parts) == 1:
        return (1, _TERMINATION_ORDER.get(term, 99))
    return (2, len(parts), _z_key(term))


_COLORS = {
    GOOD: ("C6EFCE", "006100"),
    BAD: ("FFC7CE", "9C0006"),
    MISSING: ("FFEB9C", "9C5700"),
    OUT_OF_SCOPE: ("EDEDED", "808080"),
}


def grid_axes(
    result: CheckResult | None = None,
) -> tuple[list[int], list[str | None]]:
    """Thickness indices and terminations spanned by the grid.

    Args:
        result: Output of `check_dataset`; any n or termination (including
            mixed terminations such as `O-F`) found in its files is added to
            the defaults.

    Returns:
        The sorted thickness indices and the terminations: pristine first,
        then single terminations in the order of `Termination`, then mixed
        ones.
    """
    ns = set(DEFAULT_THICKNESSES)
    terms = set(DEFAULT_TERMINATIONS)
    for rec in result.records if result else ():
        if rec.n in THICKNESSES:
            ns.add(rec.n)
        if rec.terminations:
            terms.add(rec.terminations)
    return sorted(ns), sorted(terms, key=_termination_key)


def grid_columns(
    thicknesses: Collection[int] = DEFAULT_THICKNESSES,
    terminations: Collection[str | None] = DEFAULT_TERMINATIONS,
) -> list[tuple[int, str | None, str]]:
    """All (n, termination, folder label) columns of the grid, in order.

    Args:
        thicknesses: Thickness indices to include, in order.
        terminations: Terminations to include (None = pristine), in order.

    Returns:
        One tuple per column; 50 for the defaults (n = 1-3; none, F, O).
    """
    cols = []
    for n in thicknesses:
        for term in terminations:
            for stacking in _stackings(n):
                sites = (1, 2) if term else (None,)
                for site in sites:
                    label = stacking if site is None else f"{stacking}-{site}"
                    cols.append((n, term, label))
    return cols


def grid_key(rec: CheckRecord) -> GridKey | None:
    """The grid cell a checked file belongs to, if it can be placed.

    Args:
        rec: One record of `check_dataset`.

    Returns:
        The cell key, or None if the chemistry or label could not be read.
    """
    if not rec.metals or rec.n is None:
        return None
    try:
        stacking, site, _ = parse_folder_label(rec.folderLabel)
    except ValueError:
        return None
    label = stacking if site is None else f"{stacking}-{site}"
    return (rec.metals, rec.nonmetals, rec.n, rec.terminations or None, label)


def cell_name(key: GridKey) -> str:
    """Readable name of a grid cell.

    Args:
        key: The cell key.

    Returns:
        The `mxeneId` for a single-element M, X and T, e.g. `Ti3C2O2-h1a-2`;
        otherwise a description, e.g. `Mo-Ti C n=2 O h1a-2`.
    """
    metals, nonmetals, n, term, label = key
    single = all(
        len(split_elements(part)) == 1 for part in (metals, nonmetals, term or "F")
    )
    if single:
        formula = f"{metals}{n + 1}{nonmetals}{n if n > 1 else ''}"
        return f"{formula}{term + '2' if term else ''}-{label}"
    return f"{metals} {nonmetals} n={n} {term or 'pristine'} {label}"


def systems_in(result: CheckResult) -> list[tuple[str, str]]:
    """Rows of the grid: (metals, nonmetals) systems.

    Every single metal found in the data gets a row with C and with N; every
    other system (several metals, or C-N) gets a row if found.

    Args:
        result: Output of `check_dataset`.

    Returns:
        Systems, single metals first, each group ordered by atomic number.
    """
    found = {(r.metals, r.nonmetals) for r in result.records if r.metals}
    systems = {(m, x) for m, _ in found if "-" not in m for x in NONMETALS}
    systems |= found
    return sorted(
        systems,
        key=lambda s: (len(split_elements(s[0])), _z_key(s[0]), _NONMETAL_ORDER[s[1]]),
    )


def metals_in(result: CheckResult) -> list[str]:
    """Metal systems (e.g. `Ti` or `Mo-Ti`) that appear in the checked files.

    Args:
        result: Output of `check_dataset`.

    Returns:
        Metal systems, single metals first, each group by atomic number.
    """
    return list(dict.fromkeys(m for m, _ in systems_in(result)))


def parse_scope_exclusion(text: str) -> ScopeExclusion:
    """Parse `M:T` (e.g. `Hf:O`, `Re:none`, `*:O`, `Mo-Ti:O-F`).

    Args:
        text: A metal symbol, a `-`-joined list of metals or `*`; a colon; and
            `none` or a termination, or a `-`-joined list of terminations
            (`Termination` in `schemas/labels.py`).

    Returns:
        (metals, terminations), both in canonical order (by atomic number),
        with terminations None for `none`.

    Raises:
        ValueError: If the text is malformed or names an unknown element or
            termination.
    """
    metals, sep, term = text.partition(":")
    metals, term = metals.strip(), term.strip()
    if not sep or not metals or not term:
        raise ValueError(f"Expected METAL:TERMINATION, e.g. Hf:O, got {text!r}")
    if metals != "*":
        symbols = split_elements(metals)
        bad = [s for s in symbols if not Element.is_valid_symbol(s)]
        if bad:
            raise ValueError(f"{bad} are not element symbols")
        metals = join_elements(symbols)
    if term.lower() == "none":
        return metals, None
    symbols = split_elements(term)
    bad = [s for s in symbols if s not in TERMINATIONS]
    if bad:
        raise ValueError(
            f"Termination must be none or built from "
            f"{', '.join(get_args(Termination))}, got {term!r}"
        )
    return metals, join_elements(symbols)


def _excluded(metals: str, term: str | None, out_of_scope) -> bool:
    return (metals, term) in out_of_scope or ("*", term) in out_of_scope


def grid_states(
    result: CheckResult,
    out_of_scope: Collection[ScopeExclusion] = (),
) -> tuple[dict[GridKey, str], dict[GridKey, list[CheckRecord]]]:
    """State of every grid cell and the records that claim each cell.

    Args:
        result: Output of `check_dataset`.
        out_of_scope: (metals, terminations) combinations that are not
            expected to have structures; their empty cells are shown as out of
            scope instead of missing.

    Returns:
        `{(metals, nonmetals, n, terminations, label): symbol}` and
        `{same key: check records}`.
    """
    out_of_scope = set(out_of_scope)
    claims: dict[GridKey, list[CheckRecord]] = {}
    for rec in result.records:
        key = grid_key(rec)
        if key is not None and rec.status != "skipped":
            claims.setdefault(key, []).append(rec)

    columns = grid_columns(*grid_axes(result))
    states = {}
    for m, x in systems_in(result):
        for n, term, label in columns:
            key = (m, x, n, term, label)
            recs = claims.get(key, [])
            if any(r.status == "failed" for r in recs):
                state = BAD
            elif recs:
                state = GOOD
            elif _excluded(m, term, out_of_scope):
                state = OUT_OF_SCOPE
            else:
                state = MISSING
            states[key] = state
    return states, claims


def write_overview(
    result: CheckResult,
    path: str | Path,
    root: str | Path,
    out_of_scope: Collection[ScopeExclusion] = (),
) -> None:
    """Write the overview workbook.

    Args:
        result: Output of `check_dataset`.
        path: Excel file to write.
        root: Root folder of the CONTCAR tree, for relative paths.
        out_of_scope: Combinations excluded from the scope (see `grid_states`).
    """
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    root = Path(root)
    states, claims = grid_states(result, out_of_scope)
    columns = grid_columns(*grid_axes(result))
    systems = systems_in(result)

    fills = {k: PatternFill("solid", fgColor=bg) for k, (bg, _) in _COLORS.items()}
    fonts = {k: Font(color=fg, bold=True, size=12) for k, (_, fg) in _COLORS.items()}
    center = Alignment(horizontal="center", vertical="center")
    thin, thick = Side(style="thin", color="BFBFBF"), Side(
        style="medium", color="404040"
    )
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    bold = Font(bold=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"

    # --- title, summary and legend -----------------------------------------
    totals = Counter(states.values())
    ws["A1"] = "two_d_mxenes dataset status"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        f"Generated {datetime.now().astimezone().date().isoformat()} from {root.name}: "
        f"{totals[GOOD]} good, {totals[BAD]} with problems "
        f"({len(result.failures)} files), "
        f"{totals[MISSING]} missing, {len(result.skipped)} auxiliary files skipped"
    )
    legend = [
        (GOOD, "CONTCAR passes all checks"),
        (BAD, "problem: hover the cell, or see Problems"),
        (MISSING, "no CONTCAR yet: see Missing"),
        (OUT_OF_SCOPE, "out of scope"),
    ]
    col = 1
    for symbol, text in legend:
        c = ws.cell(row=3, column=col, value=symbol or " ")
        c.fill, c.font, c.alignment = fills[symbol], fonts[symbol], center
        ws.cell(row=3, column=col + 1, value=text)
        col += 13

    # --- header rows ---------------------------------------------------------
    top, first_data = 5, 8
    first_grid_col = 2
    ws.cell(row=top + 2, column=1, value="system").font = bold
    for r in range(top, top + 3):
        ws.cell(row=r, column=1).fill = header_fill

    def group_starts(key):
        starts, prev = [], None
        for i, c in enumerate(columns):
            if key(c) != prev:
                starts.append(i)
                prev = key(c)
        return starts

    n_starts = set(group_starts(lambda c: c[0]))
    t_starts = set(group_starts(lambda c: c[:2]))

    for row, group_key, names in [
        (top, lambda c: c[0], lambda c: _n_name(c[0])),
        (top + 1, lambda c: c[:2], lambda c: _t_name(c[1])),
    ]:
        starts = sorted(group_starts(group_key)) + [len(columns)]
        for a, b in pairwise(starts):
            c1, c2 = first_grid_col + a, first_grid_col + b - 1
            ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
            cell = ws.cell(row=row, column=c1, value=names(columns[a]))
            cell.font, cell.alignment, cell.fill = bold, center, header_fill

    for i, (n, term, label) in enumerate(columns):
        cell = ws.cell(row=top + 2, column=first_grid_col + i, value=label)
        cell.alignment = Alignment(horizontal="center", text_rotation=90)
        cell.fill, cell.font = header_fill, Font(size=9)

    # --- grid ----------------------------------------------------------------
    total_col = first_grid_col + len(columns)
    for j, symbol in enumerate((GOOD, BAD, MISSING)):
        c = ws.cell(row=top + 2, column=total_col + j, value=symbol)
        c.fill, c.font, c.alignment = fills[symbol], fonts[symbol], center
    ws.cell(row=top + 1, column=total_col, value="per system").font = bold

    for r, (m, x) in enumerate(systems, start=first_data):
        ws.cell(row=r, column=1, value=f"{m} | {x}").font = bold
        row_counts: Counter[str] = Counter()
        for i, (n, term, label) in enumerate(columns):
            state = states[(m, x, n, term, label)]
            row_counts[state] += 1
            cell = ws.cell(row=r, column=first_grid_col + i, value=state or None)
            cell.fill, cell.font, cell.alignment = fills[state], fonts[state], center
            left = (
                thick
                if i in n_starts
                else (Side(style="thin", color="808080") if i in t_starts else thin)
            )
            cell.border = Border(left=left, right=thin, top=thin, bottom=thin)
            if state == BAD:
                key = (m, x, n, term, label)
                text = "\n".join(
                    f"{rec.path.relative_to(root)}: {rec.message}"
                    for rec in claims[key]
                )
                cell.comment = Comment(
                    f"{cell_name(key)}\n{text}", "dataset check", width=420, height=160
                )
        for j, symbol in enumerate((GOOD, BAD, MISSING)):
            ws.cell(row=r, column=total_col + j, value=row_counts[symbol]).alignment = (
                center
            )

    # per-column totals
    last = first_data + len(systems) - 1
    for k, symbol in enumerate((GOOD, BAD, MISSING)):
        r = last + 1 + k
        c = ws.cell(row=r, column=1, value=f"{symbol} per structure")
        c.font = Font(color=_COLORS[symbol][1], bold=True)
        for i, (n, term, label) in enumerate(columns):
            count = sum(states[(m, x, n, term, label)] == symbol for m, x in systems)
            ws.cell(row=r, column=first_grid_col + i, value=count or None).alignment = (
                center
            )

    ws.column_dimensions["A"].width = max(
        16, *(len(f"{m} | {x}") + 2 for m, x in systems)
    )
    for i in range(len(columns) + 3):
        ws.column_dimensions[get_column_letter(first_grid_col + i)].width = 4.2
    ws.row_dimensions[top + 2].height = 42
    ws.freeze_panes = ws.cell(row=first_data, column=first_grid_col)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # --- Problems -------------------------------------------------------------
    ps = wb.create_sheet("Problems")
    ps.append(["structure", "file", "measured coordination", "problem"])
    for rec in result.failures:
        ps.append(
            [
                rec.cellId or "(could not be placed)",
                str(rec.path.relative_to(root)),
                rec.measuredCoordination,
                rec.message,
            ]
        )
    _style_table(ps, [18, 45, 22, 110], bold, header_fill)

    # --- Missing ----------------------------------------------------------------
    ms = wb.create_sheet("Missing")
    ms.append(["structure", "M", "X", "n", "termination", "label"])
    for key, state in states.items():
        if state == MISSING:
            m, x, n, term, label = key
            ms.append([cell_name(key), m, x, n, term or "none", label])
    _style_table(ms, [24, 10, 6, 5, 13, 8], bold, header_fill)

    # --- All files ----------------------------------------------------------------
    fs = wb.create_sheet("All files")
    fs.append(
        ["file", "status", "structure", "formula", "measured coordination", "message"]
    )
    for rec in result.records:
        fs.append(
            [
                str(rec.path.relative_to(root)),
                rec.status,
                rec.cellId,
                rec.formula,
                rec.measuredCoordination,
                rec.message,
            ]
        )
    _style_table(fs, [50, 9, 18, 12, 22, 90], bold, header_fill)

    wb.save(path)


def _style_table(ws, widths, bold, fill) -> None:
    from openpyxl.utils import get_column_letter

    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        ws.cell(row=1, column=i).font = bold
        ws.cell(row=1, column=i).fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
